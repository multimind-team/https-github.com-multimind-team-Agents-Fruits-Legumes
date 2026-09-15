"""Validation déterministe des lignes extraites d'une facture directe.

Le modèle lit la photo/PDF et produit du JSON. Ce module refuse toute ligne
ambigüe avant l'écriture du stock ou du classeur de marge.
"""
import copy
import hashlib
import json
import math
import os
import re
import tempfile
from contextlib import contextmanager
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from threading import RLock
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.comments import Comment
import catalogue
from verrou_donnees import verrou_donnees, append_jsonl


class FactureInvalide(ValueError):
    pass


_VERROU_THREADS = RLock()


@contextmanager
def _verrou_import(racine):
    """Verrou coopératif interprocessus : lecture → dédoublonnage → deux sorties.

    Le fichier verrou est conservé (le supprimer créerait deux verrous possibles).
    Les anciens écrivains qui ne prennent pas ce verrou restent hors garantie.
    """
    with verrou_donnees(Path(racine) / "donnees"):
        yield


def _nombre(ligne, champ, *, decimal=False):
    try:
        valeur = float(ligne[champ])
        exacte = Decimal(str(ligne[champ])) if decimal else None
    except (KeyError, TypeError, ValueError, InvalidOperation):
        raise FactureInvalide(f"{champ} illisible") from None
    if not math.isfinite(valeur):
        raise FactureInvalide(f"{champ} non fini")
    if valeur < 0:
        raise FactureInvalide(f"{champ} négatif")
    return exacte if decimal else valeur


def valider_facture(donnees, correspondances):
    """Valide une facture du dataset TerreAzur, pas un mapping multi-fournisseur.

    ``correspondances`` est la table ``codes`` de codes-terreazur.json. Seuls
    TerreAzur et Pomona sont attestés pour ce dataset (_lisez_moi du mapping,
    reference/le-metier.md). La casse est indifférente ; aucun autre alias,
    découpage ou rapprochement approximatif n'est déduit. L'identité reçue
    reste conservée, notamment pour les empreintes de reprise existantes.
    """
    for champ in ("fournisseur", "date_reception", "bordereau", "pages_lues", "pages_totales", "lignes"):
        if not donnees.get(champ):
            raise FactureInvalide(f"Champ obligatoire absent : {champ}")
    fournisseur = donnees["fournisseur"]
    if not isinstance(fournisseur, str) or fournisseur.upper() not in ("TERREAZUR", "POMONA"):
        raise FactureInvalide("Fournisseur non vérifié pour le mapping TerreAzur/Pomona")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(donnees["date_reception"])):
        raise FactureInvalide("date_reception doit être AAAA-MM-JJ")
    try:
        date.fromisoformat(donnees["date_reception"])
    except ValueError:
        raise FactureInvalide("date_reception : date civile impossible") from None
    if not isinstance(donnees["lignes"], list) or not donnees["lignes"]:
        raise FactureInvalide("Aucune ligne de facture")
    if any(not re.fullmatch(r"[0-9]+", str(donnees[champ]))
           for champ in ("pages_lues", "pages_totales")):
        raise FactureInvalide("Nombre de pages illisible : entier requis, sans troncature")
    try:
        pages_lues = int(donnees["pages_lues"])
        pages_totales = int(donnees["pages_totales"])
    except (TypeError, ValueError):
        raise FactureInvalide("Nombre de pages illisible") from None
    if pages_lues < 1 or pages_lues != pages_totales:
        raise FactureInvalide("Bordereau incomplet : toutes les pages doivent être lues")

    lignes = []
    for numero, ligne in enumerate(donnees["lignes"], start=1):
        code = str(ligne.get("code_fournisseur") or "").strip()
        fiche = correspondances.get(code)
        if not fiche or not fiche.get("itm8"):
            raise FactureInvalide(f"Correspondance non vérifiée pour le code {code or 'absent'}")
        quantite = _nombre(ligne, "quantite_uf", decimal=True)
        pu = _nombre(ligne, "pu", decimal=True)
        montant = _nombre(ligne, "montant_ht", decimal=True)
        # Facultatifs : les anciens JSON restent stricts, sans inventer un PV.
        historique_pv = {}
        if "pv_magasin_ttc" in ligne:
            historique_pv["pv_magasin_ttc"] = _nombre(ligne, "pv_magasin_ttc")
            if historique_pv["pv_magasin_ttc"] == 0:
                raise FactureInvalide(f"Ligne {numero} : PV magasin historique nul")
        if "source_pv" in ligne:
            if not isinstance(ligne["source_pv"], str) or not ligne["source_pv"].strip():
                raise FactureInvalide(f"Ligne {numero} : source du PV magasin absente ou non textuelle")
            historique_pv["source_pv"] = ligne["source_pv"]
        if quantite == 0 or pu == 0:
            raise FactureInvalide(f"Ligne {numero} : quantité ou PU nul")
        # Comparer les décimaux sources sans arrondir ni élargir le centime.
        if abs((quantite * pu) - montant) > Decimal("0.01"):
            raise FactureInvalide(
                f"Ligne {numero} : Montant HT ({montant:g}) ≠ Qté fact. UF × PU ({quantite * pu:g})"
            )
        colis = ligne.get("colis")
        if colis not in (None, ""):
            colis = _nombre(ligne, "colis")
        lignes.append({
            "numero": numero,
            "code_fournisseur": code,
            "article": str(fiche["itm8"]),
            "produit": str(ligne.get("produit") or fiche.get("libelle_magasin") or "").strip(),
            "quantite": float(quantite),
            **({"unite_uf": ligne["unite_uf"]} if "unite_uf" in ligne else {}),
            "pu": float(pu),
            **({"_pv_historique": historique_pv} if historique_pv else {}),
            "montant_ht": float(montant),
            "colis": colis,
        })
    return {**donnees, "lignes": lignes}


def choisir_classeur(racine, date_reception, chemin=None):
    """Sélectionne un classeur mensuel existant, jamais un faux modèle neuf."""
    mois = date.fromisoformat(date_reception).strftime("%m%y")
    cible = (Path(chemin) if chemin is not None else
             Path(racine) / "documents-partages" / "calcul-marge-pomona" /
             f"{mois} Calcul marge Pomona.xlsx")
    if not cible.name.startswith(f"{mois} "):
        raise FactureInvalide(f"Classeur incompatible avec le mois {mois} de la facture")
    if not cible.is_file():
        raise FactureInvalide(f"Classeur du mois {mois} absent : {cible}")
    return cible


def _copier_mise_en_forme_conditionnelle(source, cible):
    for plage in source.conditional_formatting:
        for regle in source.conditional_formatting[plage]:
            cible.conditional_formatting.add(str(plage.sqref), copy.deepcopy(regle))


def _empreinte(valeur):
    contenu = json.dumps(valeur, sort_keys=True, ensure_ascii=False, allow_nan=False,
                         separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(contenu).hexdigest()


def _identite(facture):
    return {"fournisseur": str(facture["fournisseur"]).upper(),
            "bordereau": str(facture["bordereau"]), "date_reception": facture["date_reception"]}


def _ligne_empreinte(ligne):
    """Reconstitue strictement le payload v1, sans exposer de PV à Excel.

    Le déplacement en métadonnées privées ne doit pas migrer les empreintes
    existantes. Sans ces champs dans l'extraction, aucun champ n'est inventé.
    """
    return {**{cle: valeur for cle, valeur in ligne.items() if cle != "_pv_historique"},
            **ligne.get("_pv_historique", {})}


def _trace_ligne(facture, ligne):
    identite = _identite(facture)
    return {"version": 1, "facture_id": _empreinte(identite),
            "empreinte_facture": _empreinte({**identite, "lignes": [
                _ligne_empreinte(entree) for entree in facture["lignes"]]}),
            "empreinte_ligne": _empreinte(_ligne_empreinte(ligne)),
            "id": (f"livraison:{identite['date_reception']}:{ligne['article']}:direct-"
                   f"{identite['fournisseur'].lower()}:{identite['bordereau']}:{ligne['numero']}")}


def preparer_marge(chemin_classeur, jour, facture):
    """Vérifie et prépare en mémoire : aucune sauvegarde, même temporaire."""
    chemin_classeur = Path(chemin_classeur)
    try:
        with open(chemin_classeur, "rb") as source_classeur:
            classeur = load_workbook(source_classeur, data_only=False)
    except (BadZipFile, KeyError, ValueError, ParseError) as erreur:
        raise FactureInvalide(f"Le classeur de marge est illisible : {erreur}") from None
    if "Vierge" not in classeur.sheetnames:
        raise FactureInvalide("La feuille Vierge est absente du classeur de marge")
    if jour not in classeur.sheetnames:
        source = classeur["Vierge"]
        feuille = classeur.copy_worksheet(source)
        feuille.title = jour
        _copier_mise_en_forme_conditionnelle(source, feuille)
    else:
        feuille = classeur[jour]

    traces = {_trace_ligne(facture, entree)["id"]: _trace_ligne(facture, entree)
              for entree in facture["lignes"]}
    presentes = {}
    sans_trace = []
    for numero in range(4, feuille.max_row + 1):
        cellule = feuille.cell(numero, 1)
        if cellule.comment and cellule.comment.author == "Import facture directe v1":
            try:
                trace = json.loads(cellule.comment.text)
                premiere = next(iter(traces.values()))
                if trace["facture_id"] != premiere["facture_id"]:
                    continue
                if trace != traces.get(trace["id"]) or trace["id"] in presentes:
                    raise FactureInvalide("Conflit d'empreinte dans le classeur de marge")
                presentes[trace["id"]] = numero
            except (KeyError, TypeError, json.JSONDecodeError):
                raise FactureInvalide("Trace de facture illisible dans le classeur") from None
        elif cellule.value is not None:
            sans_trace.append(numero)

    manquantes = []
    for entree in facture["lignes"]:
        trace = _trace_ligne(facture, entree)
        attendues = [entree["produit"], entree["pu"], entree["quantite"]]
        numero = presentes.get(trace["id"])
        if numero is not None:
            if ([feuille.cell(numero, c).value for c in (1, 3, 6)] != attendues or
                    feuille.cell(numero, 1).data_type != "s"):
                raise FactureInvalide("Ligne de marge tracée modifiée : contrôle manuel requis")
            continue
        if any([feuille.cell(n, c).value for c in (1, 3, 6)] == attendues for n in sans_trace):
            raise FactureInvalide("Ligne de marge ancienne sans trace : rapprochement manuel requis")
        manquantes.append(entree)

    libres = [numero for numero in range(4, feuille.max_row + 1)
              if all(feuille.cell(numero, c).value is None and feuille.cell(numero, c).comment is None
                     for c in range(1, 7))]
    if len(manquantes) > len(libres):
        raise FactureInvalide("La feuille de marge ne contient pas assez de lignes libres")

    for ligne, entree in zip(libres, manquantes):
        feuille.cell(ligne, 1).value = entree["produit"]
        feuille.cell(ligne, 1).data_type = "s"
        feuille.cell(ligne, 1).comment = Comment(
            json.dumps(_trace_ligne(facture, entree), ensure_ascii=False, sort_keys=True),
            "Import facture directe v1")
        feuille.cell(ligne, 3).value = entree["pu"]
        feuille.cell(ligne, 6).value = entree["quantite"]
    classeur.calculation.fullCalcOnLoad = True
    classeur.calculation.forceFullCalc = True
    return classeur, len(manquantes)


def _stager_marge(classeur, chemin_classeur):
    """Sérialise avant les faits, puis prépare le remplacement atomique."""
    contenu = BytesIO()
    try:
        classeur.save(contenu)
    finally:
        classeur.close()
    # Détecte notamment un classeur ouvert en écriture exclusive par Excel.
    with open(chemin_classeur, "r+b"):
        pass
    temporaire = None
    try:
        with tempfile.NamedTemporaryFile(dir=Path(chemin_classeur).parent,
                                         prefix=".marge-", suffix=".xlsx", delete=False) as sortie:
            temporaire = Path(sortie.name)
            sortie.write(contenu.getvalue())
            sortie.flush()
            os.fsync(sortie.fileno())
        return temporaire
    except BaseException:
        if temporaire is not None:
            temporaire.unlink(missing_ok=True)
        raise


def ajouter_marge(chemin_classeur, jour, facture):
    """Publie le classeur sans exposer une archive ZIP partiellement écrite."""
    classeur, ajoutees = preparer_marge(chemin_classeur, jour, facture)
    if not ajoutees:
        classeur.close()
        return 0
    temporaire = _stager_marge(classeur, chemin_classeur)
    try:
        os.replace(temporaire, chemin_classeur)
    finally:
        temporaire.unlink(missing_ok=True)


def normaliser_unite(valeur):
    valeur = str(valeur or "").strip().casefold()
    return {"kg": "kg", "kilogramme": "kg", "kilogrammes": "kg",
            "pièce": "piece", "pièces": "piece", "piece": "piece", "pieces": "piece",
            "unité": "piece", "unite": "piece", "u": "piece"}.get(valeur, valeur)


def preparer_livraisons(racine, facture):
    """Vérifie le catalogue et les doublons sans modifier les faits."""
    racine = Path(racine)
    date_reception = facture["date_reception"]
    fournisseur = str(facture["fournisseur"]).upper()
    bordereau = str(facture["bordereau"])
    fichier = racine / "donnees" / "faits" / f"{date_reception[:4]}.jsonl"
    identifiants = {}
    traces = [_trace_ligne(facture, ligne) for ligne in facture["lignes"]]
    attendus = {trace["id"] for trace in traces}
    if fichier.exists():
        for texte in fichier.read_text(encoding="utf-8").splitlines():
            if not texte.strip():
                continue
            try:
                fait = json.loads(texte)
            except json.JSONDecodeError:
                raise FactureInvalide("Carnet de faits illisible : aucune écriture, récupération manuelle requise") from None
            source = fait.get("source") or {}
            meme_facture = (source.get("facture_id") == traces[0]["facture_id"] or
                           (source.get("origine") == f"mail/{fournisseur.lower()}" and
                            source.get("bordereau") == bordereau and fait.get("date_source") == date_reception))
            if meme_facture and (fait.get("id") not in attendus or
                    (source.get("empreinte_facture") is not None and
                     source["empreinte_facture"] != traces[0]["empreinte_facture"])):
                raise FactureInvalide("Conflit d'empreinte : facture déjà enregistrée avec un autre contenu")
            identifiants.setdefault(fait.get("id"), []).append(fait)

    if fichier.exists() and fichier.read_bytes()[-1:] not in (b"", b"\n"):
        raise FactureInvalide("Carnet incomplet : dernière ligne non terminée ; aucun ajout effectué")
    nouveaux = []
    for ligne in facture["lignes"]:
        fiche = catalogue.fiche(ligne["article"])
        if not fiche:
            raise FactureInvalide(f"Article absent du catalogue : {ligne['article']}")
        unite = re.sub(r"^\d+\s*", "", catalogue.unite(ligne["article"])).strip() or "inconnue"
        trace = _trace_ligne(facture, ligne)
        identifiant = trace["id"]
        mouvement = {
            "id": identifiant,
            "type": "livraison",
            "date_source": date_reception,
            "date_effet": date_reception,
            "article": ligne["article"],
            "article_source": ligne["code_fournisseur"],
            "quantite": ligne["quantite"],
            "unite": unite,
            "libelle": ligne["produit"],
            "colis": ligne["colis"],
            "par_colis": (round(ligne["quantite"] / ligne["colis"], 3)
                          if ligne["colis"] else None),
            "source": {"origine": f"mail/{fournisseur.lower()}", "bordereau": bordereau,
                       "pu": ligne["pu"], "montant_ht": ligne["montant_ht"],
                       **ligne.get("_pv_historique", {}),
                       **{cle: trace[cle] for cle in ("facture_id", "empreinte_facture", "empreinte_ligne")}},
        }
        if "unite_uf" in ligne:
            mouvement["source"]["unite_uf"] = ligne["unite_uf"]
        if identifiant in identifiants:
            for ancien in identifiants[identifiant]:
                if any(ancien.get(cle) != mouvement[cle] for cle in
                       ("type", "date_source", "date_effet", "article", "article_source",
                        "quantite", "libelle", "colis")) or any(
                        ancien["source"].get(cle) != mouvement["source"].get(cle)
                        for cle in ancien["source"]):
                    raise FactureInvalide("Conflit d'empreinte : livraison déjà enregistrée avec un autre contenu")
            continue
        unite_source = normaliser_unite(ligne.get("unite_uf"))
        unite_cible = normaliser_unite(unite)
        if not unite_source or unite_source in {"uf", "inconnue"} or unite_source != unite_cible:
            raise FactureInvalide(f"Ligne {ligne['numero']} : unité physique UF absente ou incompatible "
                                 f"avec le stock ({unite}). Relever unite_uf sur la facture ; "
                                 "aucune conversion par hypothèse.")
        nouveaux.append(mouvement)
    return fichier, nouveaux


def _publier_livraisons(fichier, nouveaux):
    if nouveaux:
        append_jsonl(fichier, nouveaux)
    return len(nouveaux)


def ecrire_livraisons(racine, facture):
    """Écrit les lignes validées dans le carnet append-only, sans doublon."""
    with _verrou_import(racine):
        return _publier_livraisons(*preparer_livraisons(racine, facture))


def importer_facture(racine, chemin_classeur, facture, simuler=False, *, classeur_seul=False):
    """Simulation sans écriture ; classeur seul sans consulter/écrire les faits.

    Le mode normal reste stock + Excel et nécessite l'autorisation métier
    habituelle. L'option classeur seul n'accorde aucun droit de stock.
    """
    if simuler:
        return _importer_facture(racine, chemin_classeur, facture, simuler=True,
                                 classeur_seul=classeur_seul)
    with _verrou_import(racine):
        return _importer_facture(racine, chemin_classeur, facture, classeur_seul=classeur_seul)


def _importer_facture(racine, chemin_classeur, facture, simuler=False, *, classeur_seul=False):
    """Publication récupérable, pas transaction atomique JSONL + Excel.

    Après interruption, rejouer le même JSON complète les faits et la marge.
    Un dernier enregistrement JSONL tronqué reste un refus manuel : jamais de
    troncature/réécriture automatique du carnet. La simulation ne réserve rien.
    """
    fichier, nouveaux = (None, []) if classeur_seul else preparer_livraisons(racine, facture)
    jour = facture["date_reception"][8:10]
    classeur, ajoutees = preparer_marge(chemin_classeur, jour, facture)
    bilan = {"livraisons_nouvelles": len(nouveaux), "feuille_marge": jour,
             "lignes_marge_nouvelles": ajoutees}
    if simuler:
        classeur.close()
        return bilan
    if not ajoutees:
        classeur.close()
        if not classeur_seul:
            _publier_livraisons(fichier, nouveaux)
        return bilan
    temporaire = _stager_marge(classeur, chemin_classeur)
    try:
        if not classeur_seul:
            _publier_livraisons(fichier, nouveaux)
        os.replace(temporaire, chemin_classeur)
    except OSError as erreur:
        if classeur_seul:
            raise FactureInvalide(
                f"Publication du classeur interrompue : {erreur}. Aucun stock écrit ; "
                "rejouer exactement le même JSON avec --classeur-seul."
            ) from erreur
        raise FactureInvalide(
            f"Publication interrompue : {erreur}. Le stock peut être déjà écrit ; "
            "la marge n'est pas confirmée. Rejouer exactement le même JSON pour "
            "reprendre sans doubler les livraisons."
        ) from erreur
    finally:
        temporaire.unlink(missing_ok=True)
    return bilan
