"""Régressions facture/marge, exclusivement dans des dossiers temporaires."""
import copy
import importlib.util
import io
import json
import os
import shutil
import subprocess
from pathlib import Path
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from threading import Event
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

MOTEUR = Path(__file__).resolve().parents[1] / "moteur"
sys.path.insert(0, str(MOTEUR))
import facture_directe as fd
import catalogue

spec = importlib.util.spec_from_file_location("importeur_facture_reprise", MOTEUR / "importer-facture-directe.py")
cli = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cli)

_WRITE_ROOT = None


def _garde_ecritures(event, args):
    if _WRITE_ROOT is None:
        return
    paths = ()
    if event == "open":
        chemin, mode, flags = args
        if ((mode and any(c in mode for c in "wax+")) or
                flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC)):
            paths = (chemin,)
    elif event in ("os.remove", "os.mkdir", "os.rmdir"):
        paths = (args[0],)
    elif event == "os.rename":
        paths = args[:2]
    for chemin in paths:
        if isinstance(chemin, (str, bytes, os.PathLike)):
            cible = Path(os.fsdecode(chemin)).resolve()
            if cible != _WRITE_ROOT and _WRITE_ROOT not in cible.parents:
                raise AssertionError(f"Test non isolé : écriture interdite {cible}")


sys.addaudithook(_garde_ecritures)


def entree():
    return {"fournisseur": "POMONA", "date_reception": "2026-10-07", "bordereau": "TEST-01",
            "pages_lues": 1, "pages_totales": 1, "lignes": [
                {"code_fournisseur": "F01", "produit": "Poireau", "quantite_uf": 12, "unite_uf": "kg",
                 "pu": 2.5, "montant_ht": 30, "colis": 2, "pv_magasin_ttc": 3.29,
                 "source_pv": "étiquette rayon 07/10/2026 validée par responsable"}]}


class FactureRepriseTests(unittest.TestCase):
    def test_unite_facture_absente_refuse_stock_mais_autorise_classeur_seul(self):
        self.donnees["lignes"][0].pop("unite_uf")
        marge = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertIn("unité physique", err)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), marge)
        code, out, err = self.appeler("--marge", self.marge, "--classeur-seul")
        self.assertEqual(code, 0, err)
        self.assertFalse(self.stock.exists())

    def test_unite_source_incompatible_refuse_avant_stock_et_marge(self):
        self.donnees["lignes"][0]["unite_uf"] = "pièce"
        marge = self.marge.read_bytes()
        code, out, err = self.appeler("--simuler")
        self.assertEqual(code, 2, out)
        self.assertIn("incompatible", err)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), marge)

    def setUp(self):
        global _WRITE_ROOT
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.fin_garde)
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        _WRITE_ROOT = self.root.resolve()
        p = patch.object(tempfile, "tempdir", str(self.root))
        p.start()
        self.addCleanup(p.stop)
        self.donnees = entree()
        self.mapping = {"F01": {"itm8": "ARTICLE01", "libelle_magasin": "Poireau"}}
        self.fichier = self.root / "facture.json"
        self.mappings = self.root / "donnees/fournisseurs/codes-terreazur.json"
        self.mappings.parent.mkdir(parents=True)
        self.mappings.write_text(json.dumps({"codes": self.mapping}), encoding="utf-8")
        self.cat = self.root / "donnees/catalogue.json"
        self.cat.write_text(json.dumps({"articles": {"ARTICLE01": {"UNITE MESURE": "KG"}}}), encoding="utf-8")
        self.marge = self.root / "documents-partages/calcul-marge-pomona/1026 Calcul marge Pomona.xlsx"
        self.creer_classeur(self.marge)
        # L'ancien CLI a un défaut global capturé par argparse : ne jamais
        # autoriser ce chemin réel pendant la phase RED.
        if hasattr(cli, "MARGE"):
            p = patch.object(cli, "MARGE", self.marge.with_name("0926 Calcul marge Pomona.xlsx"))
            p.start()
            self.addCleanup(p.stop)
        self.stock = self.root / "donnees/faits/2026.jsonl"
        for obj, attr, valeur in ((cli, "RACINE", self.root), (cli, "MAPPINGS", self.mappings),
                                  (catalogue, "FICHIER", self.cat)):
            p = patch.object(obj, attr, valeur)
            p.start()
            self.addCleanup(p.stop)
        catalogue.charger.cache_clear()
        self.addCleanup(catalogue.charger.cache_clear)

    def fin_garde(self):
        global _WRITE_ROOT
        _WRITE_ROOT = None

    def creer_classeur(self, chemin, lignes=8):
        chemin.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Vierge"
        for ligne in range(4, lignes + 1):
            ws.cell(ligne, 7, f"=C{ligne}*F{ligne}")
        wb.save(chemin)
        wb.close()

    def appeler(self, *args):
        self.fichier.write_text(json.dumps(self.donnees, ensure_ascii=False), encoding="utf-8")
        out, err = io.StringIO(), io.StringIO()
        with patch.object(sys, "argv", ["importer-facture-directe.py", str(self.fichier), *map(str, args)]), \
             redirect_stdout(out), redirect_stderr(err):
            code = cli.main()
        return code, out.getvalue(), err.getvalue()

    def test_nouvelle_facture_sans_pv_importee_sans_prix_invente_et_rejeu_noop(self):
        for champ in ("pv_magasin_ttc", "source_pv"):
            self.donnees["lignes"][0].pop(champ)
        entree_avant = copy.deepcopy(self.donnees)
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(self.donnees, entree_avant)
        fait = json.loads(self.stock.read_text(encoding="utf-8"))
        self.assertNotIn("pv_magasin_ttc", fait["source"])
        self.assertNotIn("source_pv", fait["source"])
        wb = load_workbook(self.marge)
        self.assertEqual([wb["07"].cell(4, c).value for c in range(1, 7)],
                         ["Poireau", None, 2.5, None, None, 12])
        self.assertIsNone(wb["07"]["E4"].comment)
        wb.close()
        avant = self.stock.read_bytes(), self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["livraisons_nouvelles"], 0)
        self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 0)
        self.assertEqual(avant, (self.stock.read_bytes(), self.marge.read_bytes()))

    def test_cli_copie_moteur_demain_classeur_seul_simulation_et_rejeu_sans_stock(self):
        # Vrai processus CLI, aucune racine ni dépendance métier de production.
        moteur_copie = self.root / "moteur"
        shutil.copytree(MOTEUR, moteur_copie, ignore=shutil.ignore_patterns("__pycache__"))
        demain = date.today() + timedelta(days=1)
        self.donnees["date_reception"] = demain.isoformat()
        self.donnees["bordereau"] = "FIXTURE-UNIQUEMENT-DEMAIN"
        for champ in ("pv_magasin_ttc", "source_pv"):
            self.donnees["lignes"][0].pop(champ)
        self.fichier.write_text(json.dumps(self.donnees, ensure_ascii=False), encoding="utf-8")
        cible = self.marge.with_name(f"{demain:%m%y} Calcul marge Pomona - fixture.xlsx")
        self.creer_classeur(cible)
        self.cat.unlink()  # L'Excel n'a besoin que des codes attestés, pas du stock.
        self.stock.parent.mkdir(parents=True)
        self.stock.write_bytes(b'{"id":"carnet-fixture-ne-pas-lire"')
        programme = r'''
import os, runpy, sys
from pathlib import Path
root = Path.cwd().resolve()
simulation = "--simuler" in sys.argv
def garde(event, args):
    paths = ()
    if event == "open":
        path, mode, flags = args
        if isinstance(path, (str, bytes, os.PathLike)):
            cible = Path(os.fsdecode(path)).resolve()
            if root / "donnees/faits" in cible.parents or cible == root / "donnees/catalogue.json":
                raise AssertionError("Le mode classeur seul a consulté le stock/catalogue")
        if (mode and any(c in mode for c in "wax+")) or flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC):
            paths = (path,)
    elif event in ("os.remove", "os.rmdir", "os.mkdir", "os.chmod", "os.utime"):
        paths = (args[0],)
    elif event == "os.rename":
        paths = args[:2]
    for path in paths:
        if isinstance(path, (str, bytes, os.PathLike)):
            cible = Path(os.fsdecode(path)).resolve()
            if simulation or not cible.is_relative_to(root):
                raise AssertionError(f"Écriture CLI interdite : {cible}")
sys.addaudithook(garde)
sys.path.insert(0, str(root / "moteur"))
sys.argv[0] = str(root / "moteur/importer-facture-directe.py")
runpy.run_path(sys.argv[0], run_name="__main__")
'''
        commande = [sys.executable, "-c", programme, str(self.fichier),
                    "--marge", str(cible), "--classeur-seul"]
        env = {**os.environ, "PYTHONIOENCODING": "utf-8",
               "TMPDIR": str(self.root), "TMP": str(self.root), "TEMP": str(self.root)}
        avant = {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
        simulation = subprocess.run([*commande, "--simuler"], cwd=self.root, capture_output=True,
                                    text=True, encoding="utf-8", env=env, timeout=30)
        self.assertEqual(simulation.returncode, 0, simulation.stderr)
        self.assertEqual(json.loads(simulation.stdout)["livraisons_nouvelles"], 0)
        self.assertEqual(avant, {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()})
        for passage in range(2):
            resultat = subprocess.run(commande, cwd=self.root, capture_output=True, text=True,
                                      encoding="utf-8", env=env, timeout=30)
            self.assertEqual(resultat.returncode, 0, resultat.stderr)
            bilan = json.loads(resultat.stdout)
            self.assertTrue(bilan["classeur_seul"])
            self.assertEqual(bilan["stock"], "non écrit")
            self.assertEqual(bilan["livraisons_nouvelles"], 0)
            self.assertEqual(bilan["lignes_marge_nouvelles"], 1 if passage == 0 else 0)
            if passage == 0:
                publie = cible.read_bytes()
            else:
                self.assertEqual(cible.read_bytes(), publie)
        wb = load_workbook(cible)
        ws = wb[f"{demain:%d}"]
        self.assertEqual([ws.cell(4, c).value for c in range(1, 8)],
                         ["Poireau", None, 2.5, None, None, 12, "=C4*F4"])
        self.assertIsNone(wb["Vierge"]["A4"].value)
        self.assertIsNone(ws["E4"].comment)
        wb.close()
        for relatif, contenu in avant.items():
            if self.root / relatif != cible:
                self.assertEqual((self.root / relatif).read_bytes(), contenu, relatif)
        nouveaux = {str(p.relative_to(self.root)) for p in self.root.rglob("*") if p.is_file()} - avant.keys()
        self.assertLessEqual(nouveaux, {str(Path("donnees/.operations.lock"))})

    def test_octobre_selectionne_classeur_octobre_pas_septembre(self):
        septembre = self.marge.with_name("0926 Calcul marge Pomona.xlsx")
        self.creer_classeur(septembre)
        avant = septembre.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        wb = load_workbook(self.marge)
        self.assertIn("07", wb.sheetnames)
        self.assertEqual(wb["07"]["A4"].value, "Poireau")
        wb.close()
        self.assertEqual(septembre.read_bytes(), avant)

    def test_classeur_seul_exige_une_cible_explicite(self):
        avant = self.marge.read_bytes()
        err = io.StringIO()
        with redirect_stderr(err), patch.object(sys, "argv", [
                "importer-facture-directe.py", str(self.fichier), "--classeur-seul"]):
            with self.assertRaises(SystemExit) as sortie:
                cli.main()
        self.assertEqual(sortie.exception.code, 2)
        self.assertIn("--marge", err.getvalue())
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)

    def test_mois_explicite_incompatible_est_refuse_avant_stock(self):
        septembre = self.marge.with_name("0926 Calcul marge Pomona.xlsx")
        self.creer_classeur(septembre)
        avant = septembre.read_bytes()
        code, out, err = self.appeler("--marge", septembre)
        self.assertEqual(code, 2)
        self.assertIn("mois", err.lower())
        self.assertFalse(self.stock.exists())
        self.assertEqual(septembre.read_bytes(), avant)

    def test_pv_historique_prive_empreintes_inchangees_et_colonne_e_vide(self):
        historique = copy.deepcopy(self.donnees)
        historique["lignes"][0].pop("unite_uf")  # Le payload v1 n'avait pas cette preuve physique.
        facture = fd.valider_facture(historique, self.mapping)
        ligne = facture["lignes"][0]
        self.assertNotIn("pv_magasin_ttc", ligne)
        self.assertNotIn("source_pv", ligne)
        self.assertEqual(ligne["_pv_historique"], {
            "pv_magasin_ttc": 3.29, "source_pv": self.donnees["lignes"][0]["source_pv"]})
        # Empreintes v1 enregistrées avant le déplacement en métadonnées privées.
        trace = fd._trace_ligne(facture, ligne)
        self.assertEqual(trace["facture_id"], "e491d00fc1cd112133565c2fb956a99c5c161d672eebf2fe3c523d2ddf8b2e96")
        self.assertEqual(trace["empreinte_facture"], "d2e949a3199e6b1301eb95ab08fd3900f257773dfc23374fc7035219f5f72435")
        self.assertEqual(trace["empreinte_ligne"], "799394ddae203d3e64dc13400358b1bbd65852e655cad455c4bf59bbcd54acc6")
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        fait = json.loads(self.stock.read_text(encoding="utf-8").splitlines()[0])
        self.assertEqual(fait["source"].get("source_pv"), self.donnees["lignes"][0]["source_pv"])
        self.assertEqual(fait["source"].get("pv_magasin_ttc"), 3.29)
        wb = load_workbook(self.marge)
        self.assertIsNone(wb["07"]["E4"].value)
        self.assertIsNone(wb["07"]["E4"].comment)
        wb.close()

    def test_designation_commencant_par_egal_reste_texte_litteral(self):
        self.donnees["lignes"][0]["produit"] = "=1+1"
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        wb = load_workbook(self.marge, data_only=False)
        self.assertEqual(wb["07"]["A4"].value, "=1+1")
        self.assertEqual(wb["07"]["A4"].data_type, "s")
        self.assertEqual(wb["07"]["G4"].data_type, "f")
        wb.close()

    def test_rejeu_preserve_b_d_e_historiques_et_manuels_y_compris_commentaires(self):
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        stock = self.stock.read_bytes()
        # Le premier état reconstitue E tel qu'écrit par l'ancien importeur.
        for pv, commentaire in ((3.29, self.donnees["lignes"][0]["source_pv"]),
                                (4.99, "Saisie humaine après import"), (None, None)):
            with self.subTest(pv=pv):
                wb = load_workbook(self.marge)
                ws = wb["07"]
                ws["B4"], ws["D4"], ws["E4"] = 1.23, 5.67, pv
                ws["B4"].comment = Comment("PA humain", "Responsable")
                ws["D4"].comment = Comment("Prix humain", "Responsable")
                ws["E4"].comment = Comment(commentaire, "Source PV magasin") if commentaire else None
                wb.save(self.marge)
                wb.close()
                avant = self.marge.read_bytes()
                for options in ((), ("--marge", self.marge, "--classeur-seul")):
                    code, out, err = self.appeler(*options)
                    self.assertEqual(code, 0, err)
                    self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 0)
                    self.assertEqual(self.stock.read_bytes(), stock)
                    self.assertEqual(self.marge.read_bytes(), avant)

    def test_ajout_nouvelle_facture_preserve_saisies_humaines_et_formules_existantes(self):
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        stock = self.stock.read_bytes()
        wb = load_workbook(self.marge)
        ws = wb["07"]
        for colonne, valeur in (("B", 1.25), ("D", 3.49), ("E", 4.99)):
            ws[f"{colonne}4"] = valeur
            ws[f"{colonne}4"].comment = Comment("Saisie humaine conservée", "Responsable")
        # Ligne partiellement saisie à ne jamais recycler comme ligne libre.
        ws["E5"] = 6.99
        ws["E5"].comment = Comment("Saisie en attente", "Responsable")
        wb.save(self.marge)
        wb.close()
        self.donnees["bordereau"] = "TEST-NOUVELLE-FACTURE"
        for cle in ("pv_magasin_ttc", "source_pv"):
            self.donnees["lignes"][0].pop(cle)
        code, out, err = self.appeler("--marge", self.marge, "--classeur-seul")
        self.assertEqual(code, 0, err)
        self.assertEqual(self.stock.read_bytes(), stock)
        wb = load_workbook(self.marge)
        ws = wb["07"]
        self.assertEqual([ws.cell(4, c).value for c in range(1, 8)],
                         ["Poireau", 1.25, 2.5, 3.49, 4.99, 12, "=C4*F4"])
        for colonne in ("B", "D", "E"):
            self.assertEqual(ws[f"{colonne}4"].comment.text, "Saisie humaine conservée")
        self.assertEqual(ws["E5"].value, 6.99)
        self.assertEqual(ws["E5"].comment.text, "Saisie en attente")
        self.assertIsNone(ws["A5"].value)
        self.assertEqual([ws.cell(6, c).value for c in range(1, 8)],
                         ["Poireau", None, 2.5, None, None, 12, "=C6*F6"])
        self.assertIsNone(wb["Vierge"]["A4"].value)
        self.assertEqual(wb["Vierge"]["G4"].value, "=C4*F4")
        wb.close()

    def test_conflit_cellule_a_c_f_tracee_jamais_masque_par_classeur_seul(self):
        for champ in ("pv_magasin_ttc", "source_pv"):
            self.donnees["lignes"][0].pop(champ)
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        original = self.marge.read_bytes()
        stock = self.stock.read_bytes()
        for cellule, valeur in (("A4", "Autre produit"), ("C4", 2.6), ("F4", 13)):
            for classeur_seul in (False, True):
                for simuler in (False, True):
                    with self.subTest(cellule=cellule, classeur_seul=classeur_seul, simuler=simuler):
                        self.marge.write_bytes(original)
                        wb = load_workbook(self.marge)
                        wb["07"][cellule] = valeur
                        wb.save(self.marge)
                        wb.close()
                        avant = self.marge.read_bytes()
                        options = ["--marge", self.marge]
                        if classeur_seul:
                            options.append("--classeur-seul")
                        if simuler:
                            options.append("--simuler")
                        code, out, err = self.appeler(*options)
                        self.assertEqual(code, 2, out)
                        self.assertIn("tracée modifiée", err)
                        self.assertEqual(self.marge.read_bytes(), avant)
                        self.assertEqual(self.stock.read_bytes(), stock)

    def test_rejeu_json_historique_modifie_ou_pv_retire_reste_un_conflit(self):
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        stock, marge = self.stock.read_bytes(), self.marge.read_bytes()
        changements = ({"produit": "Autre produit"}, {"pu": 3, "montant_ht": 36},
                       {"quantite_uf": 13, "montant_ht": 32.5},
                       {"source_pv": "Autre provenance"}, {"pv_magasin_ttc": 4.99}, None)
        for changement in changements:
            for classeur_seul in (False, True):
                with self.subTest(changement=changement, classeur_seul=classeur_seul):
                    self.donnees = entree()
                    if changement is None:
                        for cle in ("pv_magasin_ttc", "source_pv"):
                            self.donnees["lignes"][0].pop(cle)
                    else:
                        self.donnees["lignes"][0].update(changement)
                    options = ["--marge", self.marge]
                    if classeur_seul:
                        options.append("--classeur-seul")
                    code, out, err = self.appeler(*options)
                    self.assertEqual(code, 2, out)
                    self.assertIn("empreinte", err)
                    self.assertEqual(self.stock.read_bytes(), stock)
                    self.assertEqual(self.marge.read_bytes(), marge)

    def test_classeur_seul_sans_pv_garde_refus_codes_pages_quantites_et_montants(self):
        original = entree()
        for cle in ("pv_magasin_ttc", "source_pv"):
            original["lignes"][0].pop(cle)
        for anomalie, valeur, motif in (("code_fournisseur", "INCONNU", "Correspondance"),
                                       ("quantite_uf", 0, "nul"), ("pu", "NaN", "non fini"),
                                       ("montant_ht", "30.01001", "Montant HT"),
                                       ("pages_totales", 2, "pages")):
            with self.subTest(anomalie=anomalie):
                self.donnees = copy.deepcopy(original)
                cible = self.donnees if anomalie == "pages_totales" else self.donnees["lignes"][0]
                cible[anomalie] = valeur
                avant = self.marge.read_bytes()
                code, out, err = self.appeler("--marge", self.marge, "--classeur-seul")
                self.assertEqual(code, 2, out)
                self.assertIn(motif, err)
                self.assertEqual(self.marge.read_bytes(), avant)
                self.assertFalse(self.stock.exists())
                self.assertFalse((self.root / "donnees/.factures-directes.lock").exists())

    def test_classeur_seul_echec_publication_necrit_aucun_stock_et_rejeu_repare(self):
        avant = self.marge.read_bytes()
        with patch.object(fd.os, "replace", side_effect=PermissionError("Excel fixture verrouillé")):
            code, out, err = self.appeler("--marge", self.marge, "--classeur-seul")
        self.assertEqual(code, 2, out)
        self.assertIn("Aucun stock écrit", err)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)
        self.assertFalse(list(self.marge.parent.glob(".marge-*")))
        code, out, err = self.appeler("--marge", self.marge, "--classeur-seul")
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 1)
        self.assertFalse(self.stock.exists())

    def test_classeur_sans_vierge_refuse_avant_tout_stock(self):
        wb = load_workbook(self.marge)
        wb.active.title = "Pas de modèle"
        wb.save(self.marge)
        wb.close()
        avant = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertIn("Vierge", err)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)

    def test_simulation_verifie_catalogue_modele_et_capacite_sans_ecrire(self):
        for anomalie in ("catalogue", "Vierge", "capacité"):
            with self.subTest(anomalie=anomalie):
                self.creer_classeur(self.marge, lignes=4)
                self.cat.write_text(json.dumps({"articles": {} if anomalie == "catalogue" else
                                    {"ARTICLE01": {"UNITE MESURE": "KG"}}}), encoding="utf-8")
                catalogue.charger.cache_clear()
                if anomalie == "Vierge":
                    wb = load_workbook(self.marge)
                    wb.active.title = "Sans modèle"
                    wb.save(self.marge)
                    wb.close()
                self.donnees = entree()
                if anomalie == "capacité":
                    self.donnees["lignes"].append(copy.deepcopy(self.donnees["lignes"][0]))
                self.fichier.write_text(json.dumps(self.donnees, ensure_ascii=False), encoding="utf-8")
                avant = {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
                code, out, err = self.appeler("--simuler")
                self.assertEqual(code, 2, out)
                apres = {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
                self.assertEqual(avant, apres)

    def test_echec_serialisation_classeur_necrit_pas_stock(self):
        avant = self.marge.read_bytes()
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disque fixture")):
            code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)

    def test_reprise_apres_stock_ecrit_repare_marge_sans_doubler_puis_noop(self):
        avant = self.marge.read_bytes()
        with patch.object(fd.os, "replace", side_effect=PermissionError("Excel verrouillé fixture")):
            code, out, err = self.appeler()
        self.assertEqual(code, 2)
        self.assertIn("Rejouer exactement", err)
        self.assertEqual(self.marge.read_bytes(), avant)
        stock = self.stock.read_bytes()
        self.assertEqual(len(stock.splitlines()), 1)
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["livraisons_nouvelles"], 0)
        self.assertEqual(self.stock.read_bytes(), stock)
        wb = load_workbook(self.marge)
        self.assertEqual(wb["07"]["A4"].value, "Poireau")
        wb.close()
        repare = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(self.stock.read_bytes(), stock)
        self.assertEqual(self.marge.read_bytes(), repare)

    def test_rejeu_contenu_modifie_refuse_sans_modifier_faits_ni_marge(self):
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        stock, marge = self.stock.read_bytes(), self.marge.read_bytes()
        self.donnees["lignes"][0]["pv_magasin_ttc"] = 4.99
        code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertIn("empreinte", err)
        self.assertEqual(self.stock.read_bytes(), stock)
        self.assertEqual(self.marge.read_bytes(), marge)

    def test_empreinte_facture_et_ligne_durables_dans_faits_et_marge(self):
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        source = json.loads(self.stock.read_text(encoding="utf-8").splitlines()[0])["source"]
        for champ in ("facture_id", "empreinte_facture", "empreinte_ligne"):
            self.assertRegex(source.get(champ, ""), r"^[0-9a-f]{64}$")
        wb = load_workbook(self.marge)
        self.assertIsNotNone(wb["07"]["A4"].comment)
        trace = json.loads(wb["07"]["A4"].comment.text)
        self.assertEqual(trace["empreinte_facture"], source["empreinte_facture"])
        self.assertEqual(trace["empreinte_ligne"], source["empreinte_ligne"])
        wb.close()

    def test_reprise_apres_append_partiel_complete_seulement_lignes_manquantes(self):
        self.donnees["lignes"].append({**copy.deepcopy(self.donnees["lignes"][0]), "produit": "Poireau deuxième lot"})
        publier = fd._publier_livraisons
        def interrompre(fichier, nouveaux):
            publier(fichier, nouveaux[:1])
            raise OSError("interruption après une ligne complète")
        avant = self.marge.read_bytes()
        with patch.object(fd, "_publier_livraisons", side_effect=interrompre):
            code, out, err = self.appeler()
        self.assertEqual(code, 2)
        premier = self.stock.read_bytes()
        self.assertEqual(len(premier.splitlines()), 1)
        self.assertEqual(self.marge.read_bytes(), avant)
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["livraisons_nouvelles"], 1)
        self.assertTrue(self.stock.read_bytes().startswith(premier))
        self.assertEqual(len(self.stock.read_bytes().splitlines()), 2)
        wb = load_workbook(self.marge)
        self.assertEqual([wb["07"].cell(n, 1).value for n in (4, 5)], ["Poireau", "Poireau deuxième lot"])
        wb.close()

    def test_stock_historique_sans_empreinte_et_marge_absente_est_reparable(self):
        facture = fd.valider_facture(self.donnees, self.mapping)
        fichier, mouvements = fd.preparer_livraisons(self.root, facture)
        for cle in ("facture_id", "empreinte_facture", "empreinte_ligne", "source_pv", "pv_magasin_ttc"):
            mouvements[0]["source"].pop(cle, None)
        fd._publier_livraisons(fichier, mouvements)
        avant = self.stock.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(self.stock.read_bytes(), avant)
        self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 1)
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 0)

    def test_simulation_acceptee_ne_modifie_aucun_fichier(self):
        self.fichier.write_text(json.dumps(self.donnees, ensure_ascii=False), encoding="utf-8")
        avant = {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
        code, out, err = self.appeler("--simuler")
        self.assertEqual(code, 0, err)
        self.assertEqual(avant, {str(p.relative_to(self.root)): p.read_bytes() for p in self.root.rglob("*") if p.is_file()})

    def test_classeur_mensuel_absent_ne_cree_jamais_un_modele(self):
        self.marge.unlink()
        code, out, err = self.appeler()
        self.assertEqual(code, 2)
        self.assertIn("mois 1026 absent", err)
        self.assertFalse(self.marge.exists())
        self.assertFalse(self.stock.exists())

    def test_classeur_corrompu_est_un_refus_controle(self):
        self.marge.write_bytes(b"archive fixture invalide")
        code, out, err = self.appeler("--simuler")
        self.assertEqual(code, 2)
        self.assertIn("classeur", err.lower())
        self.assertFalse(self.stock.exists())

    def test_source_pv_non_textuelle_refusee(self):
        for valeur in ({"catalogue": "non vérifié"}, float("nan"), float("inf"),
                       float("-inf"), True, None, ""):
            with self.subTest(valeur=valeur):
                self.donnees["lignes"][0]["source_pv"] = valeur
                with self.assertRaisesRegex(fd.FactureInvalide, "source.*PV"):
                    fd.valider_facture(self.donnees, self.mapping)

    def test_append_refuse_derniere_ligne_sans_retour_sans_modifier_stock_ni_marge(self):
        self.stock.parent.mkdir(parents=True)
        prefixe = b'{"id":"historique","source":{}}'
        self.stock.write_bytes(prefixe)
        marge = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertIn("non terminée", err)
        self.assertEqual(self.stock.read_bytes(), prefixe)
        self.assertEqual(self.marge.read_bytes(), marge)

    def test_repare_uniquement_ligne_marge_disparue(self):
        self.donnees["lignes"].append({**copy.deepcopy(self.donnees["lignes"][0]), "produit": "Deuxième lot"})
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        stock = self.stock.read_bytes()
        wb = load_workbook(self.marge)
        for c in (1, 3, 5, 6):
            wb["07"].cell(5, c).value = None
            wb["07"].cell(5, c).comment = None
        wb.save(self.marge)
        wb.close()
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 1)
        self.assertEqual(self.stock.read_bytes(), stock)
        wb = load_workbook(self.marge)
        self.assertEqual([wb["07"].cell(n, 1).value for n in (4, 5, 6)], ["Poireau", "Deuxième lot", None])
        wb.close()

    def test_ligne_ancienne_sans_trace_ambigue_refusee_sans_doublon(self):
        wb = load_workbook(self.marge)
        ws = wb.copy_worksheet(wb["Vierge"])
        ws.title = "07"
        for c, valeur in ((1, "Poireau"), (3, 2.5), (5, 9.99), (6, 12)):
            ws.cell(4, c).value = valeur
        wb.save(self.marge)
        wb.close()
        avant = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2)
        self.assertIn("rapprochement manuel", err)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)

    def test_carnet_tronque_refuse_sans_effacer_la_trace(self):
        self.stock.parent.mkdir(parents=True)
        avant = b'{"id":"interruption"'
        self.stock.write_bytes(avant)
        marge = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2)
        self.assertIn("récupération manuelle", err)
        self.assertEqual(self.stock.read_bytes(), avant)
        self.assertEqual(self.marge.read_bytes(), marge)

    def test_pages_fractionnaires_ou_booleennes_ne_sont_jamais_tronquees(self):
        for valeur in (1.9, True, float("inf"), "1.5"):
            with self.subTest(valeur=valeur):
                self.donnees["pages_lues"] = valeur
                self.donnees["pages_totales"] = 1
                with self.assertRaisesRegex(fd.FactureInvalide, "pages"):
                    fd.valider_facture(self.donnees, self.mapping)

    def test_date_civile_impossible_est_refusee(self):
        self.donnees["date_reception"] = "2026-02-30"
        with self.assertRaisesRegex(fd.FactureInvalide, "date_reception"):
            fd.valider_facture(self.donnees, self.mapping)

    def test_imports_concurrents_ne_decident_pas_sur_le_meme_stock_ancien(self):
        facture = fd.valider_facture(self.donnees, self.mapping)
        premier, second, liberer = Event(), Event(), Event()
        publier = fd._publier_livraisons

        def publication_bloquee(*args):
            if premier.is_set():
                second.set()
            else:
                premier.set()
                if not liberer.wait(5):
                    raise AssertionError("Test de concurrence bloqué")
            return publier(*args)

        with patch.object(fd, "_publier_livraisons", side_effect=publication_bloquee), \
             ThreadPoolExecutor(max_workers=2) as pool:
            a = pool.submit(fd.importer_facture, self.root, self.marge, facture)
            self.assertTrue(premier.wait(5))
            b = pool.submit(fd.importer_facture, self.root, self.marge, facture)
            concurrence_avant_publication = second.wait(0.3)
            liberer.set()
            resultats = [a.result(5), b.result(5)]
        self.assertFalse(concurrence_avant_publication, "Deux imports décident avant publication du premier")
        self.assertEqual(sum(r["livraisons_nouvelles"] for r in resultats), 1)
        self.assertEqual(len(self.stock.read_bytes().splitlines()), 1)


if __name__ == "__main__":
    unittest.main()
