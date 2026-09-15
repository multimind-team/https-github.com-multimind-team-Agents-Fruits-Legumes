import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill

MOTEUR = Path(__file__).resolve().parents[1] / "moteur"
sys.path.insert(0, str(MOTEUR))

from facture_directe import FactureInvalide, ajouter_marge, ecrire_livraisons, valider_facture


MAPPINGS = {
    "104798": {"itm8": "0000087004629", "libelle_magasin": "POIREAU VRAC"},
}


def facture_valide():
    return {
        "fournisseur": "POMONA",
        "date_reception": "2026-09-07",
        "bordereau": "7835380428",
        "pages_lues": 2,
        "pages_totales": 2,
        "lignes": [{
            "code_fournisseur": "104798", "produit": "Poireau", "quantite_uf": 12, "unite_uf": "kg",
            "pu": 2.5, "montant_ht": 30, "colis": 2,
            "pv_magasin_ttc": 3.29, "source_pv": "prix rayon vérifié",
        }],
    }


class ValidationFactureTests(unittest.TestCase):
    def test_accepte_une_ligne_rapprochee_et_verifiee(self):
        facture = valider_facture(facture_valide(), MAPPINGS)
        self.assertEqual(facture["lignes"][0]["article"], "0000087004629")
        self.assertEqual(facture["lignes"][0]["quantite"], 12.0)

    def test_refuse_un_montant_qui_ne_correspond_pas_au_pu_et_a_la_quantite(self):
        donnees = facture_valide()
        donnees["lignes"][0]["montant_ht"] = 29
        with self.assertRaisesRegex(FactureInvalide, "Montant HT"):
            valider_facture(donnees, MAPPINGS)

    def test_refuse_une_valeur_non_finie(self):
        donnees = facture_valide()
        donnees["lignes"][0]["pu"] = float("nan")
        with self.assertRaisesRegex(FactureInvalide, "non fini"):
            valider_facture(donnees, MAPPINGS)

    def test_refuse_un_code_fournisseur_sans_correspondance_verifiee(self):
        donnees = facture_valide()
        donnees["lignes"][0]["code_fournisseur"] = "inconnu"
        with self.assertRaisesRegex(FactureInvalide, "Correspondance"):
            valider_facture(donnees, MAPPINGS)

    def test_refuse_un_bordereau_incomplet(self):
        donnees = facture_valide()
        donnees["pages_lues"] = 1
        with self.assertRaisesRegex(FactureInvalide, "pages"):
            valider_facture(donnees, MAPPINGS)


class MargePomonaTests(unittest.TestCase):
    def test_copie_vierge_et_remplit_uniquement_produit_pa_direct_et_quantite(self):
        with tempfile.TemporaryDirectory() as tmp:
            chemin = Path(tmp) / "marge.xlsx"
            classeur = Workbook()
            vierge = classeur.active
            vierge.title = "Vierge"
            vierge["A2"] = "Produit"
            vierge["B2"] = "PA base"
            vierge["C2"] = "PA direct"
            vierge["D2"] = "PV préco base"
            vierge["E2"] = "PV magasin"
            vierge["F2"] = "Quantité"
            vierge["G4"] = "=C4*F4"
            vierge.conditional_formatting.add("B4:E50", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="00FF00")))
            classeur.save(chemin)

            ajouter_marge(chemin, "07", valider_facture(facture_valide(), MAPPINGS))

            resultat = load_workbook(chemin, data_only=False)
            feuille = resultat["07"]
            self.assertEqual(feuille["A4"].value, "Poireau")
            self.assertEqual(feuille["C4"].value, 2.5)
            self.assertEqual(feuille["F4"].value, 12.0)
            self.assertIsNone(feuille["B4"].value)
            self.assertIsNone(feuille["D4"].value)
            self.assertIsNone(feuille["E4"].value)
            self.assertIsNone(feuille["E4"].comment)
            self.assertEqual(feuille["G4"].value, "=C4*F4")
            self.assertEqual(len(feuille.conditional_formatting), len(resultat["Vierge"].conditional_formatting))


class LivraisonDirecteTests(unittest.TestCase):
    def test_ecrit_une_livraison_traçable_et_ne_la_double_pas(self):
        with tempfile.TemporaryDirectory() as tmp:
            racine = Path(tmp)
            facture = valider_facture(facture_valide(), MAPPINGS)
            self.assertEqual(ecrire_livraisons(racine, facture), 1)
            self.assertEqual(ecrire_livraisons(racine, facture), 0)
            lignes = (racine / "donnees/faits/2026.jsonl").read_text(encoding="utf-8").splitlines()
            mouvement = json.loads(lignes[0])
            self.assertEqual(mouvement["type"], "livraison")
            self.assertEqual(mouvement["quantite"], 12.0)
            self.assertEqual(mouvement["colis"], 2.0)
            self.assertIn("7835380428", mouvement["id"])


if __name__ == "__main__":
    unittest.main()
