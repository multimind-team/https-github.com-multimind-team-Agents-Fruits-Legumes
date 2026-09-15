"""Contrat TerreAzur/Pomona et centime exact ; aucune écriture métier réelle."""
import copy
import importlib.util
import io
import json
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

RACINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RACINE / "moteur"))
import catalogue
from facture_directe import FactureInvalide, valider_facture

# Même garde que la recette isolée : aucun fallback vers le projet réel.
from verifier_factures_isolees import garde
sys.dont_write_bytecode = True
sys.addaudithook(garde)

spec = importlib.util.spec_from_file_location(
    "importeur_fournisseur_tolerance", RACINE / "moteur/importer-facture-directe.py")
cli = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cli)

MAPPING_TERREAZUR = {"TA-TEST": {"itm8": "ARTICLE-TEST", "libelle_magasin": "Produit fixture"}}


def entree():
    # La source synthétique décrit explicitement un kilogramme livré.
    return {"fournisseur": "TERREAZUR", "date_reception": "2026-10-07",
            "bordereau": "TEST-FOURNISSEUR-CENTIME", "pages_lues": 1, "pages_totales": 1,
            "lignes": [{"code_fournisseur": "TA-TEST", "produit": "Produit fixture",
                        "quantite_uf": 1, "unite_uf": "kg", "pu": 1, "montant_ht": 1, "colis": 1,
                        "pv_magasin_ttc": 2, "source_pv": "Prix rayon fixture vérifié"}]}


class FournisseurTests(unittest.TestCase):
    def test_fournisseur_etranger_ne_peut_pas_utiliser_un_code_terreazur_connu(self):
        facture = entree()
        facture["fournisseur"] = "AUTRE-FOURNISSEUR-TEST"
        with self.assertRaisesRegex(FactureInvalide, "[Ff]ournisseur.*TerreAzur"):
            valider_facture(facture, MAPPING_TERREAZUR)

    def test_seuls_les_deux_noms_attestes_sont_acceptes_sans_recrire_identite(self):
        # Sources : codes-terreazur.json/_lisez_moi et reference/le-metier.md.
        for fournisseur in ("TERREAZUR", "TerreAzur", "terreazur", "POMONA", "Pomona", "pomona"):
            with self.subTest(fournisseur=fournisseur):
                facture = entree()
                facture["fournisseur"] = fournisseur
                avant = copy.deepcopy(facture)
                resultat = valider_facture(facture, MAPPING_TERREAZUR)
                self.assertEqual(resultat["fournisseur"], fournisseur)
                self.assertEqual(resultat["lignes"][0]["article"], "ARTICLE-TEST")
                self.assertEqual(facture, avant)

    def test_pas_dalias_invente_ni_conversion_dun_fournisseur_non_textuel(self):
        for fournisseur in ("POMONA-TEST", "TERREAZUR-TEST", "Terre Azur", "POMONA/TERREAZUR",
                            " POMONA ", 42, True, ["POMONA"], {"nom": "TERREAZUR"}):
            with self.subTest(fournisseur=fournisseur):
                facture = entree()
                facture["fournisseur"] = fournisseur
                with self.assertRaisesRegex(FactureInvalide, "[Ff]ournisseur"):
                    valider_facture(facture, MAPPING_TERREAZUR)


class ToleranceCentimeTests(unittest.TestCase):
    def test_ecart_exact_dun_centime_accepte_dans_les_deux_sens(self):
        for montant in (1.01, 0.99, "1.01", "0.99"):
            with self.subTest(montant=montant):
                facture = entree()
                facture["lignes"][0]["montant_ht"] = montant
                resultat = valider_facture(facture, MAPPING_TERREAZUR)
                self.assertEqual(resultat["lignes"][0]["montant_ht"], float(montant))
                json.dumps(resultat, allow_nan=False)

    def test_produit_decimal_calcule_avant_comparaison_au_centime(self):
        facture = entree()
        facture["lignes"][0].update(quantite_uf="0.7", pu="0.1", montant_ht="0.08")
        self.assertEqual(valider_facture(facture, MAPPING_TERREAZUR)["lignes"][0]["montant_ht"], 0.08)

    def test_ecart_nul_accepte(self):
        self.assertEqual(valider_facture(entree(), MAPPING_TERREAZUR)["lignes"][0]["montant_ht"], 1)

    def test_deux_centimes_ou_juste_plus_dun_centime_refuses_sans_arrondir(self):
        for montant in (1.02, 0.98, "1.02", "0.98", "1.010001", "0.989999",
                        "1.010000000000000001", "0.989999999999999999"):
            with self.subTest(montant=montant):
                facture = entree()
                facture["lignes"][0]["montant_ht"] = montant
                with self.assertRaisesRegex(FactureInvalide, "Montant HT"):
                    valider_facture(facture, MAPPING_TERREAZUR)

    def test_controles_numeriques_existants_restent_des_refus_identifies(self):
        for champ in ("quantite_uf", "pu", "montant_ht", "pv_magasin_ttc", "colis"):
            for valeur, motif in ((None, "illisible"), ("absent", "illisible"),
                                   (float("nan"), "non fini"), ("NaN", "non fini"),
                                   (float("inf"), "non fini"), ("-Infinity", "non fini"),
                                   (-1, "négatif")):
                # Le colis absent reste facultatif, comme avant la correction.
                if champ == "colis" and valeur is None:
                    continue
                with self.subTest(champ=champ, valeur=valeur):
                    facture = entree()
                    facture["lignes"][0][champ] = valeur
                    with self.assertRaisesRegex(FactureInvalide, f"{champ} {motif}"):
                        valider_facture(facture, MAPPING_TERREAZUR)

    def test_quantite_pu_et_pv_nuls_restent_refuses(self):
        for champ in ("quantite_uf", "pu", "pv_magasin_ttc"):
            with self.subTest(champ=champ):
                facture = entree()
                facture["lignes"][0][champ] = 0
                with self.assertRaisesRegex(FactureInvalide, "nul"):
                    valider_facture(facture, MAPPING_TERREAZUR)


class FactureCliIsoleeTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name).resolve()
        self.assertNotIn(RACINE, self.root.parents)
        self.mapping = self.root / "donnees/fournisseurs/codes-terreazur.json"
        self.mapping.parent.mkdir(parents=True)
        self.mapping.write_text(json.dumps({"codes": MAPPING_TERREAZUR}), encoding="utf-8")
        self.cat = self.root / "donnees/catalogue.json"
        self.cat.write_text(json.dumps({"articles": {"ARTICLE-TEST": {"UNITE MESURE": "KG"}}}),
                            encoding="utf-8")
        self.marge = self.root / "1026 Calcul marge Pomona.xlsx"
        wb = Workbook()
        wb.active.title = "Vierge"
        wb.active["G4"] = "=C4*F4"
        wb.save(self.marge)
        wb.close()
        self.stock = self.root / "donnees/faits/2026.jsonl"
        self.fichier = self.root / "facture.json"
        self.donnees = entree()
        for objet, nom, valeur in ((cli, "RACINE", self.root), (cli, "MAPPINGS", self.mapping),
                                    (catalogue, "FICHIER", self.cat),
                                    (tempfile, "tempdir", str(self.root))):
            p = patch.object(objet, nom, valeur)
            p.start()
            self.addCleanup(p.stop)
        catalogue.charger.cache_clear()
        self.addCleanup(catalogue.charger.cache_clear)

    def appeler(self, simuler=False):
        self.assertEqual(cli.RACINE, self.root)
        self.assertEqual(cli.MAPPINGS, self.mapping)
        self.assertEqual(catalogue.FICHIER, self.cat)
        self.assertTrue(self.marge.is_relative_to(self.root))
        self.fichier.write_text(json.dumps(self.donnees), encoding="utf-8")
        # Toujours une racine ET un --marge temporaires, même pour un refus attendu.
        args = [str(cli.__file__), str(self.fichier), "--marge", str(self.marge)]
        if simuler:
            args.append("--simuler")
        out, err = io.StringIO(), io.StringIO()
        with patch.object(sys, "argv", args), redirect_stdout(out), redirect_stderr(err):
            code = cli.main()
        return code, out.getvalue(), err.getvalue()

    def test_cli_fournisseur_etranger_refuse_avant_stock_et_marge(self):
        self.donnees["fournisseur"] = "AUTRE-FOURNISSEUR-TEST"
        avant = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertRegex(err, "[Ff]ournisseur.*TerreAzur")
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)
        self.assertFalse((self.root / "donnees/.operations.lock").exists())

    def test_cli_plus_dun_centime_refuse_avant_stock_et_marge(self):
        self.donnees["lignes"][0]["montant_ht"] = "1.02"
        avant = self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 2, out)
        self.assertIn("Montant HT", err)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)
        self.assertFalse((self.root / "donnees/.operations.lock").exists())

    def test_cli_centime_accepte_publie_et_se_rejoue_uniquement_en_fixture(self):
        self.donnees["lignes"][0]["montant_ht"] = "1.01"
        avant = self.marge.read_bytes()
        code, out, err = self.appeler(simuler=True)
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["livraisons_nouvelles"], 1)
        self.assertFalse(self.stock.exists())
        self.assertEqual(self.marge.read_bytes(), avant)
        self.assertFalse((self.root / "donnees/.operations.lock").exists())

        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        faits = [json.loads(ligne) for ligne in self.stock.read_text(encoding="utf-8").splitlines()]
        self.assertEqual(len(faits), 1)
        self.assertEqual(faits[0]["article"], "ARTICLE-TEST")
        self.assertEqual(faits[0]["source"]["montant_ht"], 1.01)
        self.assertEqual(faits[0]["source"]["unite_uf"], "kg")
        self.assertEqual(faits[0]["source"]["source_pv"], self.donnees["lignes"][0]["source_pv"])
        wb = load_workbook(self.marge)
        try:
            self.assertEqual([wb["07"].cell(4, c).value for c in (1, 3, 5, 6)],
                             ["Produit fixture", 1, None, 1])
            self.assertIsNone(wb["07"]["E4"].comment)
        finally:
            wb.close()
        stock, marge = self.stock.read_bytes(), self.marge.read_bytes()
        code, out, err = self.appeler()
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads(out)["livraisons_nouvelles"], 0)
        self.assertEqual(json.loads(out)["lignes_marge_nouvelles"], 0)
        self.assertEqual(self.stock.read_bytes(), stock)
        self.assertEqual(self.marge.read_bytes(), marge)


if __name__ == "__main__":
    unittest.main()
