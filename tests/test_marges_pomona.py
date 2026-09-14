"""Accès indépendant du mail : vrai HTTP, classeurs temporaires uniquement."""
import http.client
import importlib.util
import json
import tempfile
import threading
import unittest
from pathlib import Path
from urllib.parse import quote
from unittest.mock import patch

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
NOM = "0926 Calcul marge Pomona - livraison 09-09-2026.xlsx"
DOSSIER = "documents-partages/calcul-marge-pomona/"


class MargesPomonaTests(unittest.TestCase):
    def setUp(self):
        spec = importlib.util.spec_from_file_location("serveur_marges_test", ROOT / "moteur/serveur.py")
        self.module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.module)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.module.RACINE = self.root
        self.module.DONNEES = self.root / "donnees"
        self.folder = self.root / DOSSIER
        self.folder.mkdir(parents=True)
        self.classeur(NOM, rempli=True)
        self.classeur("0926 Calcul marge Pomona.xlsx", rempli=False)
        self.classeur("notes-privees.xlsx", rempli=True)
        self.server = self.module.ServeurHTTP(("127.0.0.1", 0), self.module.Gestionnaire)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self.fermer)

    def fermer(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(5)

    def classeur(self, nom, rempli):
        wb = Workbook()
        wb.active.title = "Vierge"
        if rempli:
            ws = wb.create_sheet("09")
            ws.append(["ENTETE TEST"])
            ws.cell(4, 1, "POIREAU TEST")
            ws.cell(4, 3, 2.5)
            ws.cell(4, 5, 3)
            ws.cell(4, 6, 10)
        wb.save(self.folder / nom)
        wb.close()

    def requete(self, path, methode="GET", headers=None):
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=5)
        try:
            connection.request(methode, path, headers=headers or {})
            response = connection.getresponse()
            return response.status, response.read(), {k.title(): v for k, v in response.getheaders()}
        finally:
            connection.close()

    def test_liste_et_telechargement_binaire_sans_mail_ni_modele_vide(self):
        statut, corps, entetes = self.requete("/api/marges-pomona")
        self.assertEqual(statut, 200)
        self.assertEqual(entetes.get("Cache-Control"), "no-store")
        fichiers = json.loads(corps)["fichiers"]
        self.assertEqual([f["nom"] for f in fichiers], [NOM])
        self.assertEqual(fichiers[0]["jours"], ["2026-09-09"])
        statut, corps, entetes = self.requete(fichiers[0]["url"])
        self.assertEqual(statut, 200)
        self.assertEqual(corps, (self.folder / NOM).read_bytes())
        self.assertEqual(entetes.get("Cache-Control"), "no-store")
        self.assertTrue(entetes["Content-Disposition"].startswith("attachment;"))
        self.assertIn("spreadsheetml.sheet", entetes["Content-Type"])
        statut, corps, _ = self.requete(fichiers[0]["url"], "HEAD")
        self.assertEqual(statut, 200)
        self.assertEqual(corps, b"")
        self.assertFalse((self.root / "donnees").exists())

    def test_liste_reflete_nouveau_classeur_et_signale_illisible(self):
        self.classeur("1026 Calcul marge Pomona.xlsx", rempli=True)
        (self.folder / "1126 Calcul marge Pomona.xlsx").write_bytes(b"CLASSEUR CASSE TEST")
        statut, corps, _ = self.requete("/api/marges-pomona?actualisation=1")
        self.assertEqual(statut, 200)
        resultat = json.loads(corps)
        self.assertEqual([f["nom"] for f in resultat["fichiers"]], ["1026 Calcul marge Pomona.xlsx", NOM])
        self.assertEqual(resultat["fichiers"][0]["jours"], ["2026-10-09"])
        self.assertEqual(resultat["indisponibles"], 1)

    def test_pas_de_listing_prive_traversee_ou_hote_etranger(self):
        chemins = ["/" + DOSSIER.rstrip("/"), "/" + DOSSIER,
                   "/" + DOSSIER + quote("0926 Calcul marge Pomona.xlsx"),
                   "/" + DOSSIER + "notes-privees.xlsx",
                   "/" + DOSSIER + "../AGENT.md",
                   "/" + DOSSIER + "%2e%2e/AGENT.md",
                   "/" + DOSSIER + "%252e%252e/AGENT.md",
                   "/" + DOSSIER + quote(NOM) + "%00"]
        for methode in ("GET", "HEAD"):
            for chemin in chemins:
                with self.subTest(methode=methode, chemin=chemin):
                    self.assertEqual(self.requete(chemin, methode)[0], 404)
            self.assertEqual(self.requete("/api/marges-pomona", methode, {"Host": "etranger.test"})[0], 403)
        self.assertEqual(self.requete("/api/marges-pomona", "POST")[0], 403)

    def test_disparition_entre_liste_et_telechargement_reste_un_refus_http(self):
        liste = self.module.marges_pomona_disponibles()
        (self.folder / NOM).unlink()
        with patch.object(self.module, "marges_pomona_disponibles", return_value=liste):
            self.assertEqual(self.requete(liste["fichiers"][0]["url"])[0], 404)


if __name__ == "__main__":
    unittest.main()
